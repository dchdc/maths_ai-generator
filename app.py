import os, json, re, sqlite3, uuid
from pathlib import Path

import streamlit as st
import sympy as sp
from google import genai
from google.genai import types

APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / "quizzes.db"
MODEL = os.getenv("GEMINI_MODEL", "gemini-3.6-flash")

TOPICS = [
    "Whole numbers", "Fractions", "Decimals", "Percentages", "Ratio", "Speed",
    "Circles", "Nets and 3D shapes", "Heuristics", "Prime factorisation", "HCF and LCM",
    "Logarithms and exponential functions", "Algebra", "Angles and bearings", "Trigonometry",
    "Circle theorems", "Vectors", "Surds", "Polynomials", "Binomial expansion",
    "Coordinate geometry", "Differentiation", "Integration", "Other / mixed"
]

st.set_page_config(page_title="Singapore Maths AI Generator", page_icon="📘", layout="wide")


def db_init():
    with sqlite3.connect(DB_PATH) as con:
        con.execute("CREATE TABLE IF NOT EXISTS quizzes (id TEXT PRIMARY KEY, title TEXT, payload TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)")


def get_quiz(qid):
    with sqlite3.connect(DB_PATH) as con:
        row = con.execute("SELECT title, payload FROM quizzes WHERE id=?", (qid,)).fetchone()
    if not row:
        return None
    try:
        return {"title": row[0], **json.loads(row[1])}
    except Exception:
        return None


def save_quiz(title, questions, settings):
    qid = uuid.uuid4().hex[:8]
    payload = {"questions": questions, "settings": settings}
    with sqlite3.connect(DB_PATH) as con:
        con.execute("INSERT INTO quizzes(id,title,payload) VALUES(?,?,?)", (qid, title, json.dumps(payload, ensure_ascii=False)))
    return qid


def extract_file(upload):
    name = upload.name.lower()
    raw = upload.getvalue()
    text = ""
    images = []
    if name.endswith(".pdf"):
        import fitz
        doc = fitz.open(stream=raw, filetype="pdf")
        pages_text = []
        for page in doc:
            pages_text.append(page.get_text("text"))
        text = "\n\n".join(pages_text).strip()
        if len(text) < 80 and len(doc):
            # Send the first page as a visual fallback; advise user if the document is multi-page.
            pix = doc[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            images.append((pix.tobytes("png"), "image/png"))
    elif name.endswith(".docx"):
        from docx import Document
        import io
        doc = Document(io.BytesIO(raw))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(chunks)
    elif name.endswith((".png", ".jpg", ".jpeg")):
        mime = "image/png" if name.endswith(".png") else "image/jpeg"
        images.append((raw, mime))
    elif name.endswith(".txt"):
        text = raw.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type.")
    return text, images


def call_gemini(prompt, images=None):
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY is not set. Add it to your Windows environment variables, then restart the terminal.")
    client = genai.Client(api_key=api_key)
    contents = [prompt]
    for data, mime in (images or []):
        contents.append(types.Part.from_bytes(data=data, mime_type=mime))
    response = client.models.generate_content(model=MODEL, contents=contents, config=types.GenerateContentConfig(temperature=0.2))
    result = (response.text or "").strip()
    result = re.sub(r"^```(?:json)?\s*", "", result, flags=re.I)
    result = re.sub(r"\s*```$", "", result)
    return json.loads(result)


def normalize_question(q, i):
    q = dict(q) if isinstance(q, dict) else {}
    q.setdefault("question", f"Question {i}")
    q.setdefault("topic", "Other / mixed")
    q.setdefault("difficulty", "Medium")
    q.setdefault("type", "typed")
    q.setdefault("options", [])
    q.setdefault("answer", "")
    q.setdefault("worked_solution", "")
    q.setdefault("diagram_spec", None)
    q.setdefault("marks", 1)
    if q.get("type") not in ("typed", "mcq"):
        q["type"] = "typed"
    if not isinstance(q.get("options"), list):
        q["options"] = []
    try:
        q["marks"] = max(1, int(q["marks"]))
    except Exception:
        q["marks"] = 1
    return q


def equivalent(user, answer):
    a, b = str(user).strip(), str(answer).strip()
    if not a:
        return False
    if a.casefold() == b.casefold():
        return True
    try:
        return sp.simplify(sp.sympify(a.replace("^", "**")) - sp.sympify(b.replace("^", "**"))) == 0
    except Exception:
        return False


def make_prompt(source_text, topics, level, difficulty, count, answer_mode, extra, include_diagrams=True):
    return f'''You are a careful Singapore mathematics worksheet assistant. Return ONLY valid JSON, no markdown.
Create a teacher-reviewable quiz based on the uploaded worksheet content. Preserve the mathematics and use Singapore school conventions. Do not copy long passages verbatim; create similar but original questions where requested.
Student level: {level}. Requested topics: {topics}. Difficulty: {difficulty}. Number of NEW questions: {count}. Answer mode: {answer_mode}.
Teacher instructions: {extra or 'None'}
Diagram-based questions requested: {'YES — include at least 3 diagram-based questions where suitable' if include_diagrams and count >= 3 else ('YES where suitable' if include_diagrams else 'NO')}
Uploaded worksheet text (may be incomplete):
{source_text[:18000] if source_text else '[No selectable text; inspect attached page image.]'}
JSON shape:
{{"title":"short quiz title","detected_topics":["topic"],"extracted_questions":[{{"question":"...","topic":"...","difficulty":"Easy|Medium|Hard","type":"typed|mcq","options":["A. ...","B. ...","C. ...","D. ..."],"answer":"exact answer or option letter","worked_solution":"step-by-step solution","marks":2,"diagram_spec":null}}],"new_questions":[same object shape]}}
Rules: create exactly {count} new_questions. Extract as many clearly readable source questions as practical (up to 20); if source text is empty, extracted_questions may be empty. Use type mcq only if answer mode is MCQ or Both; otherwise typed. For MCQ include 4 plausible options and set answer to the correct option letter (A/B/C/D). For typed use a concise answer suitable for automatic checking. Show a mathematically sound worked solution. Do not invent a claim that an extracted question was in the source if unreadable. When diagram-based questions are requested, include at least 3 of the new questions with a useful simple diagram_spec. diagram_spec must be null for non-diagram questions, or an object such as {{"type":"circle","labels":["A","B","C"],"caption":"not to scale"}}, {{"type":"triangle","labels":["A","B","C"],"caption":"AB = 5 cm"}}, {{"type":"rectangle","labels":["8 cm","x cm"],"caption":"rectangle"}}, {{"type":"number_line","labels":["0","1","2","3"],"caption":"number line"}}, or {{"type":"coordinate_grid","points":[[1,2],[3,4]],"labels":["A","B"],"caption":"coordinate plane"}}. Keep diagram specs simple and ensure the question is answerable from the diagram and text. Ensure each question is appropriate for the level and requested topics.''' 


def build_worksheet_pdf(title, questions, include_answers=False):
    """Build a printable worksheet PDF, drawing simple schematic diagrams where specified."""
    import io, math
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Flowable
    from xml.sax.saxutils import escape

    class DiagramFlowable(Flowable):
        def __init__(self, spec):
            super().__init__(); self.spec = spec or {}; self.width = 420; self.height = 125
        def wrap(self, aw, ah):
            self.width = min(420, aw); return self.width, self.height
        def draw(self):
            c = self.canv; typ = str(self.spec.get('type','')).lower(); w=self.width; h=self.height
            c.saveState(); c.setStrokeColor(colors.HexColor('#263746')); c.setFillColor(colors.HexColor('#263746')); c.setLineWidth(1.3)
            if typ == 'circle':
                cx,cy=w/2,70; r=36; c.circle(cx,cy,r,stroke=1,fill=0)
                labels=self.spec.get('labels',[])
                positions=[(cx,cy+r+5),(cx+r+4,cy),(cx,cy-r-13)]
                for lab,pos in zip(labels,positions): c.drawString(pos[0],pos[1],str(lab))
            elif typ == 'triangle':
                pts=[(w/2,102),(w/2-48,25),(w/2+48,25)]; path=c.beginPath(); path.moveTo(*pts[0]); path.lineTo(*pts[1]); path.lineTo(*pts[2]); path.close(); c.drawPath(path)
                for lab,pos in zip(self.spec.get('labels',[]),[(w/2,107),(w/2-65,15),(w/2+50,15)]): c.drawString(pos[0],pos[1],str(lab))
            elif typ == 'rectangle':
                x,y=w/2-75,35; c.rect(x,y,150,55,stroke=1,fill=0); labs=self.spec.get('labels',[])
                if labs: c.drawCentredString(w/2,95,str(labs[0]))
                if len(labs)>1: c.drawString(x+155,y+22,str(labs[1]))
            elif typ == 'number_line':
                x1,x2=35,w-35; y=62; c.line(x1,y,x2,y); c.line(x2,y,x2-7,y+4); c.line(x2,y,x2-7,y-4)
                labs=self.spec.get('labels',[]) or ['0','1','2','3','4']; n=max(1,len(labs)-1)
                for i,lab in enumerate(labs):
                    x=x1+(x2-x1)*i/n; c.line(x,y-5,x,y+5); c.drawCentredString(x,y-18,str(lab))
            elif typ == 'coordinate_grid':
                x0,y0=w/2-85,25; scale=22
                for i in range(9): c.line(x0+i*scale,y0,x0+i*scale,y0+88); c.line(x0,y0+i*scale,x0+176,y0+i*scale)
                c.line(x0,y0+44,x0+190,y0+44); c.line(x0+88,y0,x0+88,y0+100)
                pts=self.spec.get('points',[]); labs=self.spec.get('labels',[])
                for i,pt in enumerate(pts):
                    try:
                        x=x0+88+float(pt[0])*scale; y=y0+44+float(pt[1])*scale; c.circle(x,y,3,fill=1)
                        if i<len(labs): c.drawString(x+4,y+4,str(labs[i]))
                    except Exception: pass
            else:
                c.drawString(5,65,'Diagram: ' + str(self.spec.get('caption','refer to question')))
            cap=self.spec.get('caption')
            if cap and typ in ('circle','triangle','rectangle','number_line','coordinate_grid'):
                c.setFont('Helvetica-Oblique',8); c.drawCentredString(w/2,5,str(cap)[:100])
            c.restoreState()

    buf=io.BytesIO()
    doc=SimpleDocTemplate(buf,pagesize=A4,rightMargin=1.8*cm,leftMargin=1.8*cm,topMargin=1.6*cm,bottomMargin=1.6*cm)
    styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name='WorksheetTitle',parent=styles['Title'],alignment=TA_CENTER,fontSize=17,leading=21,spaceAfter=8))
    styles.add(ParagraphStyle(name='QuestionText',parent=styles['BodyText'],fontSize=10.5,leading=15,spaceAfter=5))
    story=[Paragraph(escape(title),styles['WorksheetTitle']),Paragraph('Name: ____________________________________    Class: ______________    Date: ______________',styles['BodyText']),Spacer(1,12)]
    for i,q in enumerate(questions,1):
        story.append(Paragraph(f'<b>Question {i}</b> &nbsp; ({int(q.get("marks",1))} mark(s)) &nbsp; <font color="#536878">{escape(str(q.get("topic","")))}</font>',styles['BodyText']))
        story.append(Paragraph(escape(str(q.get('question',''))).replace('\n','<br/>'),styles['QuestionText']))
        if q.get('type')=='mcq' and q.get('options'):
            for opt in q['options']:
                story.append(Paragraph(escape(str(opt)),styles['QuestionText']))
        spec=q.get('diagram_spec')
        if isinstance(spec,dict) and spec.get('type'):
            story.append(DiagramFlowable(spec)); story.append(Spacer(1,4))
        story.append(Paragraph('Working: __________________________________________________________________________',styles['BodyText']))
        story.append(Paragraph('____________________________________________________________________________________',styles['BodyText']))
        story.append(Spacer(1,13))
    if include_answers:
        story.append(PageBreak()); story.append(Paragraph('Answer Key and Worked Solutions',styles['WorksheetTitle']))
        for i,q in enumerate(questions,1):
            story.append(Paragraph(f'<b>Q{i}. Answer:</b> {escape(str(q.get("answer","")))}',styles['BodyText']))
            story.append(Paragraph(escape(str(q.get('worked_solution','Solution not provided.'))).replace('\n','<br/>'),styles['QuestionText']))
            story.append(Spacer(1,7))
    doc.build(story); return buf.getvalue()


def student_page(qid):
    quiz = get_quiz(qid)
    if not quiz:
        st.error("Quiz not found. Check that the quiz ID is correct and the server has the saved quiz database.")
        st.stop()
    st.title(quiz.get("title", "Maths quiz"))
    st.caption("Singapore Maths practice · Answers are marked automatically")
    qs = quiz.get("questions", [])
    if not qs:
        st.warning("This quiz has no questions.")
        st.stop()
    with st.form("student_quiz"):
        responses = []
        for i, q in enumerate(qs):
            st.markdown(f"**Q{i+1}. ({q.get('marks', 1)} mark{'s' if q.get('marks', 1) != 1 else ''})** {q.get('question','')}")
            if q.get("type") == "mcq" and q.get("options"):
                responses.append(st.radio("Choose an answer", q["options"], key=f"stu_{qid}_{i}", index=None, label_visibility="collapsed"))
            else:
                responses.append(st.text_input("Your answer", key=f"stu_{qid}_{i}", label_visibility="collapsed"))
        submitted = st.form_submit_button("Submit answers", type="primary")
    if submitted:
        score = 0
        total = 0
        for i, (q, response) in enumerate(zip(qs, responses)):
            marks = int(q.get("marks", 1)); total += marks
            answer = str(q.get("answer", "")).strip()
            if q.get("type") == "mcq":
                correct = str(response or "").strip().upper().startswith(answer.upper()) if answer.upper() in ("A", "B", "C", "D") else str(response or "").strip().casefold() == answer.casefold()
            else:
                correct = equivalent(response or "", answer)
            if correct: score += marks
            with st.expander(f"Q{i+1}: {'Correct' if correct else 'Review'} — {marks if correct else 0}/{marks}"):
                st.write("**Correct answer:**", answer)
                st.write("**Worked solution:**", q.get("worked_solution", "Solution not provided."))
        st.success(f"Your score: {score}/{total} ({round(100*score/total) if total else 0}%)")
    st.stop()


def main():
    db_init()
    params = st.query_params
    qid = params.get("quiz")
    if qid:
        student_page(str(qid))
    st.title("📘 Singapore Maths AI Question Generator")
    st.write("Upload a worksheet, review what the AI extracts, generate similar practice questions, and publish a quiz link for students.")
    st.info(f"Gemini model: `{MODEL}`. AI-generated questions and answers should be checked by a teacher before publishing.")

    with st.sidebar:
        st.header("Quiz settings")
        level = st.selectbox("Student level", ["Primary 1", "Primary 2", "Primary 3", "Primary 4", "Primary 5", "Primary 6", "Secondary 1", "Secondary 2", "Secondary 3", "Secondary 4", "Secondary 5"])
        topics = st.multiselect("Topics", TOPICS, default=["Fractions"])
        difficulty = st.select_slider("Difficulty", options=["Easy", "Medium", "Hard"], value="Medium")
        count = st.slider("Number of new questions", 1, 20, 10)
        include_diagrams = st.checkbox("Include diagram-based questions where suitable", value=True)
        answer_mode = st.radio("Student answer format", ["Typed", "MCQ", "Both"], index=2)
        extra = st.text_area("Extra instructions (optional)", placeholder="e.g. Use word problems; avoid calculators; include 2-step questions")

    upload = st.file_uploader("Upload worksheet (PDF, DOCX, PNG, JPG, or TXT)", type=["pdf", "docx", "png", "jpg", "jpeg", "txt"])
    typed_source = st.text_area("Or paste worksheet text / describe the question style", height=140, placeholder="Paste questions here, or describe the topic and style to generate from scratch.")
    if st.button("Extract and generate questions", type="primary", use_container_width=True):
        try:
            source_text, images = ("", [])
            if upload:
                source_text, images = extract_file(upload)
            if typed_source.strip():
                source_text = (source_text + "\n\n" + typed_source.strip()).strip()
            if not source_text and not images:
                st.error("Upload a worksheet or paste some text first.")
                st.stop()
            prompt = make_prompt(source_text, topics or ["Other / mixed"], level, difficulty, count, answer_mode, extra, include_diagrams)
            with st.spinner(f"Asking {MODEL} to read the worksheet and create questions..."):
                data = call_gemini(prompt, images)
            st.session_state["generated_data"] = data
            st.session_state["generation_settings"] = {"level": level, "topics": topics, "difficulty": difficulty, "answer_mode": answer_mode, "include_diagrams": include_diagrams}
            st.success("Draft created. Review and edit the questions below before publishing.")
        except Exception as e:
            st.error(f"Could not generate questions: {e}")
            st.caption("Check GEMINI_API_KEY, internet access, and model availability. If Gemini returns a temporary 503/504, wait a little and try again.")

    data = st.session_state.get("generated_data")
    if data:
        st.subheader("Detected topics")
        detected = data.get("detected_topics", [])
        st.write(", ".join(map(str, detected)) if detected else "No topics detected")
        extracted = [normalize_question(q, i+1) for i, q in enumerate(data.get("extracted_questions", []))]
        generated = [normalize_question(q, i+1) for i, q in enumerate(data.get("new_questions", []))]
        if extracted:
            with st.expander(f"Questions extracted from worksheet ({len(extracted)})", expanded=False):
                for i, q in enumerate(extracted):
                    st.markdown(f"**Q{i+1}. {q['question']}**")
                    st.caption(f"Topic: {q['topic']} · Answer: {q['answer']}")
                    st.write(q["worked_solution"])
        st.subheader("Review generated questions")
        reviewed = []
        for i, q in enumerate(generated):
            with st.expander(f"Question {i+1} · {q['topic']} · {q['difficulty']}", expanded=(i == 0)):
                question = st.text_area("Question", q["question"], key=f"q_text_{i}")
                c1, c2 = st.columns(2)
                with c1:
                    topic = st.selectbox("Topic", TOPICS, index=TOPICS.index(q["topic"]) if q["topic"] in TOPICS else len(TOPICS)-1, key=f"q_topic_{i}")
                    qtype = st.selectbox("Answer type", ["typed", "mcq"], index=0 if q["type"] == "typed" else 1, key=f"q_type_{i}")
                with c2:
                    marks = st.number_input("Marks", 1, 20, int(q["marks"]), key=f"q_marks_{i}")
                    qdiff = st.selectbox("Difficulty", ["Easy", "Medium", "Hard"], index=["Easy", "Medium", "Hard"].index(q["difficulty"]) if q["difficulty"] in ["Easy", "Medium", "Hard"] else 1, key=f"q_diff_{i}")
                options_text = st.text_area("MCQ options (one per line)", "\n".join(q["options"]), key=f"q_options_{i}") if qtype == "mcq" else ""
                answer = st.text_input("Correct answer (MCQ: A/B/C/D or exact option text)", str(q["answer"]), key=f"q_answer_{i}")
                solution = st.text_area("Worked solution", str(q["worked_solution"]), key=f"q_solution_{i}")
                reviewed.append({"question": question, "topic": topic, "difficulty": qdiff, "type": qtype, "options": [x.strip() for x in options_text.splitlines() if x.strip()] if qtype == "mcq" else [], "answer": answer, "worked_solution": solution, "marks": int(marks)})
        title = st.text_input("Quiz title", data.get("title", f"{st.session_state.get('generation_settings', {}).get('level','Maths')} Practice Quiz"))
        st.subheader("Download printable worksheet")
        pdf_answers = st.checkbox("Include answer key and worked solutions in the PDF", value=False)
        try:
            pdf_bytes = build_worksheet_pdf(title or "Singapore Maths Practice", reviewed, include_answers=pdf_answers)
            st.download_button("Download questions as PDF", data=pdf_bytes, file_name="singapore_maths_worksheet.pdf", mime="application/pdf", use_container_width=True)
        except Exception as e:
            st.error(f"Could not create PDF: {e}")

        if st.button("Publish quiz and create student link", type="primary"):
            if not reviewed:
                st.error("There are no generated questions to publish.")
            elif any(not q["question"].strip() or not q["answer"].strip() for q in reviewed):
                st.error("Every question needs both a question and a correct answer.")
            elif any(q["type"] == "mcq" and len(q["options"]) < 2 for q in reviewed):
                st.error("Each MCQ needs at least two options.")
            else:
                new_id = save_quiz(title, reviewed, st.session_state.get("generation_settings", {}))
                base = str(st.get_option("server.address") or "localhost")
                port = st.get_option("server.port") or 8501
                link = f"http://{base}:{port}/?quiz={new_id}"
                st.success("Quiz saved on this server.")
                st.code(link)
                st.caption("This link works on your computer while the app is running. For students outside your network, deploy the app to a hosting service and share the hosted URL with the same ?quiz=ID suffix.")
                st.button("Copy link text", on_click=lambda: st.session_state.update({"copy_link_text": link}))
                if st.session_state.get("copy_link_text") == link:
                    st.write("Copy the link shown above and share it with students.")

    st.divider()
    st.caption("Supported extraction: selectable PDF text, DOCX paragraphs/tables, images, and pasted text. For scanned multi-page PDFs, export pages as images or use OCR before uploading; the current fallback sends only the first page image. Auto-marking is exact for MCQ and numeric/simple symbolic expressions; it is not a full exam mark scheme.")

if __name__ == "__main__":
    main()
